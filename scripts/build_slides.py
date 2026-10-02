"""Build the ME2 V3 submission deck from slides_data.json.

  python scripts/build_slides.py            # -> slides/me2_v3_submission.pptx

Reproduces the two required template slides (the Model/Dataset/Training/
Validation overview, then "To Be Submitted" + reviewer checklist) plus a small
appendix. Known values are ink; pending/estimate values are amber so the
remaining fill-ins are obvious.  Needs:  pip install python-pptx
"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / "slides_data.json").read_text(encoding="utf-8"))

# ---- palette ----------------------------------------------------------------
INK = RGBColor(0x1A, 0x1A, 0x1A)
HEAD = RGBColor(0x1F, 0x4E, 0x79)
NAVY = RGBColor(0x16, 0x23, 0x3F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x55, 0x55, 0x55)
PENDING = RGBColor(0xB8, 0x6A, 0x00)
TAN = RGBColor(0xE8, 0xC7, 0x7E)
TAN_DK = RGBColor(0x8A, 0x5A, 0x1E)
OK = RGBColor(0x1F, 0x7A, 0x3D)
GREEN = RGBColor(0x7A, 0xB8, 0x4A)
ALT = RGBColor(0xF2, 0xF6, 0xFB)
BOX = {
    "blue":   (RGBColor(0x2E, 0x5B, 0x9F), RGBColor(0xE8, 0xF0, 0xFA)),
    "green":  (RGBColor(0x3E, 0x8E, 0x5B), RGBColor(0xE6, 0xF4, 0xEA)),
    "orange": (RGBColor(0xC8, 0x8A, 0x2E), RGBColor(0xFA, 0xF0, 0xDE)),
    "purple": (RGBColor(0x7A, 0x5C, 0xA8), RGBColor(0xEF, 0xE9, 0xF6)),
}


def add_text(slide, x, y, w, h, text, size=11, bold=False, color=INK,
             align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return tb


def _setcell(cell, text, size=9.5, bold=False, color=INK, align=PP_ALIGN.LEFT):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    cell.text_frame.word_wrap = True
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"


def add_box(slide, x, y, w, h, label, text, style="blue", size=10.5,
            pending=False):
    line, fill = BOX[style]
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.09)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = label + "  "
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = line
    r2 = p.add_run()
    r2.text = text
    r2.font.size = Pt(size)
    r2.font.color.rgb = PENDING if pending else INK
    return shp


def kv_table(slide, x, y, w, items, vals, item_w, val_w, row_h,
             header=("Item", "Value"), val_status=None):
    n = len(items)
    gfx = slide.shapes.add_table(n + 1, 2, Inches(x), Inches(y),
                                 Inches(w), Inches((n + 1) * row_h))
    t = gfx.table
    t.columns[0].width = Inches(item_w)
    t.columns[1].width = Inches(val_w)
    for i in range(n + 1):
        t.rows[i].height = Inches(row_h)
    for c, htxt in enumerate(header):
        cell = t.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        _setcell(cell, htxt, size=9.5, bold=True, color=WHITE)
    for i in range(n):
        fill = WHITE if i % 2 == 0 else ALT
        c0 = t.cell(i + 1, 0)
        c0.fill.solid()
        c0.fill.fore_color.rgb = fill
        _setcell(c0, items[i], size=9.5, bold=True, color=INK)
        c1 = t.cell(i + 1, 1)
        c1.fill.solid()
        c1.fill.fore_color.rgb = fill
        st = val_status[i] if val_status else "known"
        _setcell(c1, vals[i], size=9.5,
                 color=PENDING if st in ("pending", "estimate") else INK)
    return t


def build_slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    m = DATA["meta"]
    add_text(s, 0.42, 0.12, 12.5, 0.5, m["title"], 26, True, INK)
    add_text(s, 0.42, 0.62, 12.5, 0.34, m["subtitle"], 13, True, INK)
    add_text(s, 0.42, 0.98, 12.5, 0.26, m["date_line"], 11, False, GRAY)

    model = DATA["model"]
    b = model["boxes"]
    add_text(s, 0.42, 1.30, 3.0, 0.3, "Model", 15, True, HEAD)
    add_box(s, 0.42, 1.62, 6.05, 0.34, b[0]["label"], b[0]["text"], b[0]["style"])
    add_box(s, 0.42, 2.00, 6.05, 0.38, b[1]["label"], b[1]["text"], b[1]["style"])
    add_box(s, 0.42, 2.44, 2.98, 0.32, b[2]["label"], b[2]["text"], b[2]["style"], 10)
    add_box(s, 3.49, 2.44, 2.98, 0.32, b[3]["label"], b[3]["text"], b[3]["style"], 10)
    add_box(s, 0.42, 2.80, 2.98, 0.32, b[4]["label"], b[4]["text"], b[4]["style"], 10)
    add_box(s, 3.49, 2.80, 2.98, 0.32, b[5]["label"], b[5]["text"], b[5]["style"], 10)
    kv_table(s, 0.42, 3.20, 6.05, [model["param_item"]], [model["param_value"]],
             2.5, 3.55, 0.3)

    add_text(s, 0.42, 3.92, 6.05, 0.28, "Validation on the Raspberry Pi 4",
             13, True, HEAD)
    v = DATA["validation"]["rows"]
    kv_table(s, 0.42, 4.24, 6.05, [r["item"] for r in v], [r["value"] for r in v],
             2.5, 3.55, 0.47, val_status=[r["status"] for r in v])

    ds = DATA["dataset"]["rows"]
    add_text(s, 6.95, 1.30, 3.0, 0.3, "Dataset", 15, True, HEAD)
    kv_table(s, 6.95, 1.62, 6.05, [r["item"] for r in ds], [r["value"] for r in ds],
             1.9, 4.15, 0.55, val_status=[r["status"] for r in ds])

    tr = DATA["training"]["rows"]
    add_text(s, 6.95, 4.50, 5.0, 0.28, "Training on the A100 cluster",
             13, True, HEAD)
    kv_table(s, 6.95, 4.82, 6.05, [r["item"] for r in tr], [r["value"] for r in tr],
             1.9, 4.15, 0.42, val_status=[r["status"] for r in tr])

    add_text(s, 0.42, 7.06, 12.5, 0.4,
             "Architecture note: fixed-window BC-ResNet KWS (two ONNX: wake + "
             "command/slot, shared 80-mel) \u2014 not a streaming / KV-cache encoder. "
             "Amber = pending (post-training).",
             8.5, False, GRAY, italic=True)
    return s


def build_slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(s, 0.42, 0.28, 8.0, 0.6, "To Be Submitted", 28, True, INK)
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(0.94),
                              Inches(12.5), Inches(0.05))
    rule.fill.solid()
    rule.fill.fore_color.rgb = GREEN
    rule.line.fill.background()
    rule.shadow.inherit = False

    sub = DATA["submission"]
    rows = sub["rows"]
    gfx = s.shapes.add_table(len(rows), 2, Inches(0.42), Inches(1.14),
                             Inches(12.5), Inches(len(rows) * 0.56))
    t = gfx.table
    t.columns[0].width = Inches(3.4)
    t.columns[1].width = Inches(9.1)
    for i, r in enumerate(rows):
        t.rows[i].height = Inches(0.56)
        c0 = t.cell(i, 0)
        c0.fill.solid()
        c0.fill.fore_color.rgb = TAN
        _setcell(c0, r["item"], size=11, bold=True, color=TAN_DK)
        c1 = t.cell(i, 1)
        c1.fill.solid()
        c1.fill.fore_color.rgb = WHITE
        pend = r["status"] in ("pending", "estimate")
        _setcell(c1, r["value"], size=10.5, color=PENDING if pend else INK)

    add_text(s, 0.42, 3.62, 8.0, 0.3, "What the reviewer will look for",
             15, True, HEAD)
    rev = sub["reviewer"]["rows"]
    n = len(rev)
    gfx2 = s.shapes.add_table(n + 1, 3, Inches(0.42), Inches(3.98),
                              Inches(12.5), Inches((n + 1) * 0.42))
    t2 = gfx2.table
    t2.columns[0].width = Inches(0.7)
    t2.columns[1].width = Inches(8.9)
    t2.columns[2].width = Inches(2.9)
    for i in range(n + 1):
        t2.rows[i].height = Inches(0.42)
    for c, htxt in enumerate(("#", "Item", "Status")):
        cell = t2.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        _setcell(cell, htxt, size=10, bold=True, color=WHITE)
    for i, r in enumerate(rev):
        fill = WHITE if i % 2 == 0 else ALT
        c0 = t2.cell(i + 1, 0)
        c0.fill.solid()
        c0.fill.fore_color.rgb = fill
        _setcell(c0, str(i + 1), size=9.5, color=INK, align=PP_ALIGN.CENTER)
        c1 = t2.cell(i + 1, 1)
        c1.fill.solid()
        c1.fill.fore_color.rgb = fill
        _setcell(c1, r["item"], size=9.5, color=INK)
        c2 = t2.cell(i + 1, 2)
        c2.fill.solid()
        c2.fill.fore_color.rgb = fill
        st = r["status"].lower()
        scol = OK if st == "yes" else (PENDING if "pend" in st or st == "planned"
                                       else HEAD)
        _setcell(c2, r["status"], size=9.5, bold=True, color=scol)
    return s


def build_slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(s, 0.42, 0.28, 12.5, 0.6,
             "Appendix \u2014 filling the pending numbers", 24, True, INK)
    add_text(s, 0.42, 0.95, 12.5, 0.3,
             "Amber fields on slides 1\u20132 are filled post-training, in this order:",
             12, False, GRAY)
    steps = [
        "1.  tools/manifest_stats.py        \u2192  dataset Hours / Speakers   (HPC, data/manifests/)",
        "2.  runs/v3r1/results.json         \u2192  Steps / loss, Keyword + intent acc, False-accept rate",
        "3.  Pi live benchmark (vcm-benchmark) \u2192  Latency p95 / RTF / runtime threads",
        "4.  tools/export_onnx.py --verify  \u2192  ONNX + model-weights release URL",
        "5.  edit slides_data.json, then:    python scripts/build_slides.py",
        "6.  update SUBMISSION.md statuses   \u2192  commit \u2192 push",
    ]
    y = 1.45
    for st in steps:
        add_text(s, 0.7, y, 12.0, 0.4, st, 13, False, INK)
        y += 0.46
    add_text(s, 0.42, y + 0.15, 12.5, 0.3, "Data & citations", 13, True, HEAD)
    y += 0.55
    cites = [
        "Dataset: HF airimonda/ai231-me2-voice-commands (default config; per-source licenses; research + education)",
        "Multi-Sensor Voice Command (Xela Ubalde): CC BY 4.0, DOI 10.48804/IEKKVZ",
        "Live harness: github.com/airimonda/vcm-benchmark        Code: github.com/jblagana/ai231-me2-v3 (MIT)",
    ]
    for c in cites:
        add_text(s, 0.7, y, 12.0, 0.35, c, 11.5, False, INK)
        y += 0.4
    return s


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    build_slide1(prs)
    build_slide2(prs)
    build_slide3(prs)
    out = ROOT / "slides" / "me2_v3_submission.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print("wrote", out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
